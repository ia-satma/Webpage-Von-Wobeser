#!/usr/bin/env python3
"""Build the versioned, bilingual Dropbox 2026 news snapshot.

This script is intentionally deterministic. It converts the eleven approved
Dropbox source notes to the restricted CMS HTML vocabulary, records provenance
hashes, validates both PDFs and optionally copies them to the mirror's public
PDF directory. It never connects to PostgreSQL.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl


PUBLIC_PDF_PREFIX = "/images/PDF_news/2026"
URL_RE = re.compile(r"https?://[^\s<>]+", re.IGNORECASE)
SPACE_RE = re.compile(r"\s+")
UNSAFE_HTML_RE = re.compile(
    r"</?(?:script|iframe|object|embed)\b|\bon\w+\s*=|\bjavascript\s*:",
    re.IGNORECASE,
)
STOP_PREFIXES = (
    "para obtener información adicional",
    "para obtener informacion adicional",
    "for additional information",
    "for more information, please contact",
    "para más información, favor de contactar",
    "para mas informacion, favor de contactar",
    "la información contenida en esta nota",
    "la informacion contenida en esta nota",
    "la información incluida en esta nota",
    "la informacion incluida en esta nota",
    "the information contained in this note",
    "a t e n t a m e n t e",
    "s i n c e r e l y",
    "atentamente",
    "sincerely",
)


@dataclass(frozen=True)
class SourceConfig:
    path: str
    excerpt_ordinal: int


@dataclass(frozen=True)
class NoteConfig:
    key: str
    folder: str
    date: str
    slug: str
    title: str
    title_es: str
    en: SourceConfig | None
    es: SourceConfig | None
    pdf: str
    pdf_es: str
    authors: tuple[tuple[str, str], ...]
    tags: tuple[str, ...]
    excerpt: str | None = None
    excerpt_es: str | None = None
    manual_en: tuple[str, ...] = ()
    manual_es: tuple[str, ...] = ()


NOTES: tuple[NoteConfig, ...] = (
    NoteConfig(
        key="2026-06-22-legacy-permit-migration",
        folder="Junio/22 Junio",
        date="2026-06-22",
        slug="migracion-voluntaria-permisos-legados-lspee-lse",
        title="New guidelines for the voluntary migration of legacy permits to the current electricity regulatory framework enter into force",
        title_es="Entran en vigor los lineamientos para la migración voluntaria de permisos legados al régimen eléctrico vigente",
        en=SourceConfig("Nota_ Acuerdo migración expedita LSPEE LSE (inglés).docx", 1),
        es=SourceConfig("Nota_ Acuerdo migración expedita LSPEE LSE (español).docx", 1),
        pdf="2026-06-22-Expedited-Migration-Agreement-LSPEE-LSE.pdf",
        pdf_es="2026-06-22-Acuerdo-migracion-expedita-LSPEE-LSE.pdf",
        authors=(
            ("Edmond Grieger", "egrieger@vwys.com.mx"),
            ("Ariel Garfio", "agarfio@vwys.com.mx"),
            ("Edmundo Berumen", "eberumen@vwys.com.mx"),
            ("Roberto Flores", "rflores@vwys.com.mx"),
            ("Héctor Sánchez", "hsanchez@vwys.com.mx"),
            ("Mauricio Puebla", "mpuebla@vwys.com.mx"),
            ("Regina González", "rgonzalez@vwys.com.mx"),
        ),
        tags=("energy-natural-resources", "electricity", "regulatory"),
        excerpt="On June 18, 2026, guidelines issued by the Ministry of Energy entered into force to regulate the voluntary and expedited migration of legacy self-supply and cogeneration permits to the framework of the Electric Sector Law.",
        excerpt_es="El 18 de junio de 2026 entraron en vigor los lineamientos emitidos por la Secretaría de Energía para regular la migración voluntaria y expedita de permisos legados de autoabastecimiento y cogeneración al marco de la Ley del Sector Eléctrico.",
    ),
    NoteConfig(
        key="2026-06-26-fifa-administrative-measures",
        folder="Junio/26 Junio",
        date="2026-06-26",
        slug="medidas-administrativas-copa-mundial-fifa-2026-cdmx",
        title="Administrative measures – Events related to the FIFA World Cup 2026 in Mexico City",
        title_es="Medidas administrativas – Eventos relativos a la Copa Mundial FIFA 2026 en la Ciudad de México",
        en=SourceConfig("VWYS - VWYS - Flash Informativo - Administrative Measures FIFA World Cup 2026 3 - (724466v1).docx", 2),
        es=SourceConfig("VWYS - VWYS - Flash Informativo - Mundial 3 - (724467v2).docx", 2),
        pdf="2026-06-26-Events-Related-to-the-WC-FIFA-in-Mexico-City.pdf",
        pdf_es="2026-06-26-Eventos-Relativos-con-la-copa-mundial-FIFA-en-ciudad-de-mexico.pdf",
        authors=(
            ("Rafael Vallejo", "rvallejo@vwys.com.mx"),
            ("Adrián Castillo", "adcastillo@vwys.com.mx"),
            ("Alejandra Arizpe", "aarizpe@vwys.com.mx"),
            ("Sarah Gibert", "sgibert@vwys.com.mx"),
            ("Santiago Torres", "storres@vwys.com.mx"),
            ("Ana Ruiz", "aruiz@vwys.com.mx"),
            ("Alejandro Pérez", "alperez@vwys.com.mx"),
            ("Ricardo Rosas", "rrosasg@vwys.com.mx"),
            ("Gabriela Zambrano", "gzambrano@vwys.com.mx"),
        ),
        tags=("labor-employment", "fifa-world-cup-2026", "remote-work"),
    ),
    NoteConfig(
        key="2026-07-03-ofac-fincen-cjng-fuel-smuggling",
        folder="Julio/03 Julio",
        date="2026-07-03",
        slug="ofac-fincen-contrabando-combustible-cjng",
        title="U.S. targets criminal facilitators behind CJNG’s cross-border fuel smuggling schemes: What Mexican corporations and foreign companies operating in Mexico need to know",
        title_es="EE.UU. designa a facilitadores financieros del contrabando transfronterizo de combustible del CJNG: implicaciones para empresas mexicanas y compañías extranjeras que operan en México",
        en=SourceConfig("030726 Client Alert OFAC July 2026 v2 (1).docx", 3),
        es=SourceConfig("0307026 Client_Alert_OFAC_CJNG_Julio_2026_ESP_v1 (1).docx", 3),
        pdf="2026-07-03-2026-07-03-Client-Alert-OFAC-ENG.pdf",
        pdf_es="2026-07-03-Client-Alert-OFAC-ESP.pdf",
        authors=(),
        tags=("investigations-anti-corruption-compliance", "ofac", "fincen", "energy"),
    ),
    NoteConfig(
        key="2026-07-16-tax-conclusive-agreement",
        folder="Julio/16 Julio",
        date="2026-07-16",
        slug="contingencia-fiscal-acuerdo-conclusivo",
        title="Multimillion-peso tax contingency favorably resolved through a conclusive agreement",
        title_es="Contingencia fiscal multimillonaria resuelta favorablemente mediante acuerdo conclusivo",
        en=SourceConfig("Nota_Informativa_-_Acuerdo_Conclusivo_Exitoso_140726_EN.docx", 2),
        es=SourceConfig("Nota_Informativa_-_Acuerdo_Conclusivo_Exitoso_140726_ES.docx", 2),
        pdf="2026-07-16-Multimillion-Peso-Tax-Contingency-Favorably-Resolved.pdf",
        pdf_es="2026-07-16-Contingencia-Fiscal-Millonaria-Resuelta.pdf",
        authors=(("Alejandro Torres", "ajtorres@vwys.com.mx"), ("Luis Enrique Torres", "ltorres@vwys.com.mx")),
        tags=("tax", "tax-controversy", "prodecon"),
    ),
    NoteConfig(
        key="2026-07-17-us-cartel-terrorist-designations",
        folder="Julio/17 Julio",
        date="2026-07-17",
        slug="estados-unidos-designaciones-terroristas-carteles",
        title="United States expands terrorist designations of cartels: Implications for companies in Chihuahua and Michoacán",
        title_es="Estados Unidos amplía las designaciones terroristas de cárteles: implicaciones para las empresas en Chihuahua y Michoacán",
        en=SourceConfig("170726 FTO_Designation_Client_Alert_July_2026_English.docx", 3),
        es=SourceConfig("170726 FTO_Designation_Client_Alert_Julio_2026_Español.docx", 3),
        pdf="2026-07-17-US-Expands-Terrorist-Designations-of-Cartels.pdf",
        pdf_es="2026-07-17-EEUU-Amplia-Designaciones-Terroristas-de-Carteles.pdf",
        authors=(
            ("Diego Sierra", "dsierra@vwys.com.mx"),
            ("Ricardo Cacho", "rcacho@vwys.com.mx"),
            ("Enrique Riquelme", "eriquelme@vwys.com.mx"),
            ("María Elisa Vera Madrigal", "mvera@vwys.com.mx"),
            ("Ana Victoria Guevara", "aguevara@vwys.com.mx"),
        ),
        tags=("investigations-anti-corruption-compliance", "sanctions", "foreign-terrorist-organizations"),
    ),
    NoteConfig(
        key="2026-07-21-imss-electronic-signature",
        folder="Julio/21 Julio",
        date="2026-07-21",
        slug="efirma-unico-certificado-imss",
        title="Implementation of the e.signature as the sole valid certificate before the Mexican Social Security Institute (IMSS)",
        title_es="Implementación de la e.firma como único certificado válido ante el Instituto Mexicano del Seguro Social (IMSS)",
        en=SourceConfig("Flash Informativo e.firma IMSS (ENG) (1).docx", 2),
        es=SourceConfig("Flash Informativo e.firma IMSS (ESP) (1).docx", 2),
        pdf="2026-07-21-Implementation-e.signature-IMSS.pdf",
        pdf_es="2026-07-21-Implementacion-e.firma-IMSS.pdf",
        authors=(
            ("Rafael Vallejo", "rvallejo@vwys.com.mx"),
            ("Adrián Castillo", "adcastillo@vwys.com.mx"),
            ("Alejandra Arizpe", "aarizpe@vwys.com.mx"),
            ("Sarah Gibert", "sgibert@vwys.com.mx"),
            ("Santiago Torres", "storres@vwys.com.mx"),
            ("Alejandro Pérez", "alperez@vwys.com.mx"),
            ("Ana Ruiz", "aruiz@vwys.com.mx"),
            ("Ricardo Rosas", "rrosasg@vwys.com.mx"),
        ),
        tags=("labor-employment", "imss", "electronic-signature"),
    ),
    NoteConfig(
        key="2026-07-28-impi-pct-authority",
        folder="Julio/28 Julio",
        date="2026-07-28",
        slug="impi-administracion-busqueda-internacional",
        title="The Mexican Institute of Industrial Property (IMPI) designated as an International Searching Authority and International Preliminary Examining Authority under the Patent Cooperation Treaty (PCT)",
        title_es="El Instituto Mexicano de la Propiedad Industrial (IMPI) recibe nombramiento como Administración encargada de la Búsqueda Internacional y del Examen Preliminar Internacional de patentes",
        en=None,
        es=None,
        pdf="2026-07-28-IMPI-Designated-International-Searching-Authority.pdf",
        pdf_es="2026-07-28-IMPI-Designado-Busqueda-Internacional-Patentes.pdf",
        authors=(("Patricia Kaim", "pkaim@vwys.com.mx"), ("Efrén Sánchez", "efsanchez@vwys.com.mx")),
        tags=("intellectual-property", "patents", "impi"),
        manual_en=(
            "In July 2026, the Mexican Institute of Industrial Property (IMPI) was officially appointed as an International Searching Authority (ISA) and International Preliminary Examining Authority (IPEA) under the Patent Cooperation Treaty (PCT) administered by the World Intellectual Property Organization (WIPO).",
            "With this designation, IMPI becomes the 26th International Authority worldwide and the fourth PCT International Authority operating in Spanish, joining the patent offices of Spain, Chile, and Brazil.",
            "As a result of this appointment, inventors and applicants filing international patent applications under the PCT may now designate IMPI to conduct, in Spanish, both the International Search (ISA) and the International Preliminary Examination (IPEA) for applications proceeding under Chapter I and/or Chapter II of the PCT.",
            "To support its new role, IMPI has expanded its team of substantive patent examiners and strengthened its technological infrastructure, enhancing its capacity to provide these international services.",
            "At Von Wobeser y Sierra, we remain at your disposal to advise you on this and any other patent-related matters, and to assist you in assessing how this development may benefit your international patent filing strategy.",
        ),
        manual_es=(
            "En julio de 2026, el Instituto Mexicano de la Propiedad Industrial (IMPI) fue designado oficialmente como el órgano de Administración encargado de la Búsqueda Internacional y del Examen Preliminar Internacional (ISA/IPEA), bajo el Tratado de Cooperación en materia de Patentes (PCT) de la Organización Mundial de la Propiedad Intelectual (OMPI).",
            "Con esta designación, el IMPI se convierte en la autoridad número 26 en el mundo, siendo la cuarta oficina que opera en idioma español, junto con España, Chile y Brasil.",
            "A partir de esta designación, las personas inventoras o creadoras podrán elegir al IMPI para llevar a cabo, en idioma español, tanto una búsqueda internacional (ISA) del estado de la técnica como un examen preliminar internacional (IPEA) de patentabilidad, en beneficio de aquellos solicitantes que decidan proceder bajo los Capítulos I y/o II del PCT.",
            "Con motivo de esta designación, el IMPI incrementó su número de Especialistas en Examen de Fondo y, con ello, fortaleció su infraestructura tecnológica.",
            "En Von Wobeser y Sierra quedamos a su disposición para asesorarlos en este y cualquier otro asunto relacionado con patentes, así como para ayudarles a evaluar cómo este desarrollo puede beneficiar su estrategia internacional de protección y tramitación de patentes.",
        ),
    ),
    NoteConfig(
        key="2026-07-29-cne-self-consumption-registration",
        folder="Julio/29 Julio",
        date="2026-07-29",
        slug="programa-temporal-registro-autoconsumo-petroliferos",
        title="CNE announces the implementation of a temporary registration program for the dispatch for self-consumption of petroleum products",
        title_es="La CNE anuncia la implementación de un programa temporal de registro para el despacho para autoconsumo de petrolíferos",
        en=SourceConfig("VWYS - Nota Legal- Programa temporal de Registro CNE en inglés (1).docx", 1),
        es=SourceConfig("VWYS - Nota Legal- Programa temporal de Registro CNE - (750929v1).docx", 1),
        pdf="2026-07-29-CNE-Temporary-Registration-Program-Self-Consumption.pdf",
        pdf_es="2026-07-29-CNE-Programa-Temporal-Registro-Autoconsumo-Petroliferos.pdf",
        authors=(
            ("Ariel Garfio", "agarfio@vwys.com.mx"),
            ("Edmundo Berumen", "eberumen@vwys.com.mx"),
            ("Mauricio Puebla", "mpuebla@vwys.com.mx"),
            ("Regina González", "rgonzalez@vwys.com.mx"),
            ("Arturo Hernández", "ahernandez@vwys.com.mx"),
        ),
        tags=("energy-natural-resources", "hydrocarbons", "self-consumption"),
    ),
    NoteConfig(
        key="2026-08-04-mining-protected-natural-areas",
        folder="Agosto/4 Agosto",
        date="2026-08-04",
        slug="mineria-areas-naturales-protegidas",
        title="New provisions for the evaluation and approval of mining projects in protected natural areas",
        title_es="Nuevas disposiciones para la evaluación y autorización de proyectos de minería en áreas naturales protegidas",
        en=SourceConfig("Nota - Acuerdo Minería ANP V.P. ENG (002).docx", 1),
        es=SourceConfig("Nota - Acuerdo Minería ANP V.SLG 2.docx", 1),
        pdf="2026-08-04-New-Provisions-Mining-Projects-Protected-Natural-Areas.pdf",
        pdf_es="2026-08-04-Nuevas-Disposiciones-Mineria-Areas-Naturales-Protegidas.pdf",
        authors=(),
        tags=("environmental", "mining", "protected-natural-areas"),
        excerpt="On July 20, 2026, the Ministry of the Environment and Natural Resources published an agreement that establishes significant restrictions on mining activities within federally administered protected natural areas and reinforces the prohibitions introduced by the May 2023 Mining Reform.",
        excerpt_es="El 20 de julio de 2026, la Secretaría de Medio Ambiente y Recursos Naturales publicó un acuerdo que establece restricciones significativas para actividades mineras dentro de áreas naturales protegidas federales y refuerza las prohibiciones de la Reforma Minera de mayo de 2023.",
    ),
    NoteConfig(
        key="2026-08-14-anti-money-laundering-rules",
        folder="Agosto/14 Agosto",
        date="2026-08-14",
        slug="reforma-reglas-ley-antilavado-2026",
        title="Amendment to the General Rules of the Anti-Money Laundering Law (LFPIORPI)",
        title_es="Reforma a las Reglas de Carácter General de la Ley Antilavado (LFPIORPI)",
        en=SourceConfig("VWYS - Client Alert _ Reforma RGC PLD - (inglés).docx", 1),
        es=SourceConfig("VWYS - Client Alert _ Reforma RGC PLD - (738787v10).docx", 1),
        pdf="2026-08-14-Amendment-General-Rules-Anti-Money-Laundering-Law.pdf",
        pdf_es="2026-08-14-Reforma-Reglas-Caracter-General-Ley-Antilavado.pdf",
        authors=(
            ("Luis Burgueño", "lburgueno@vwys.com.mx"),
            ("Alberto Córdoba", "acordoba@vwys.com.mx"),
            ("Ricardo Cacho", "rcacho@vwys.com.mx"),
            ("Max Morales", "mmorales@vwys.com.mx"),
            ("Joel Domínguez", "jdominguez@vwys.com.mx"),
            ("María Elisa Vera Madrigal", "mvera@vwys.com.mx"),
        ),
        tags=("investigations-anti-corruption-compliance", "anti-money-laundering", "lfpiorpi"),
    ),
    NoteConfig(
        key="2026-08-14-fracking-recommendations",
        folder="Agosto/14 Agosto Bis",
        date="2026-08-14",
        slug="recomendaciones-fracking-mexico",
        title="Scientific committee presents recommendations on the feasibility of fracking in Mexico",
        title_es="Comité científico presenta recomendaciones sobre la viabilidad del fracking en México",
        en=SourceConfig("VWYS - Nota Legal- Fracking en inglés (1).docx", 1),
        es=SourceConfig("VWYS - Nota Legal- Fracking (VWYS!765502.1) 1 (1) (1).docx", 1),
        pdf="2026-08-14-Scientific-Committee-Recommendations-Fracking-Mexico.pdf",
        pdf_es="2026-08-14-Comite-Cientifico-Recomendaciones-Fracking-Mexico.pdf",
        authors=(
            ("Ariel Garfio", "agarfio@vwys.com.mx"),
            ("Edmundo Berumen", "eberumen@vwys.com.mx"),
            ("Mauricio Puebla", "mpuebla@vwys.com.mx"),
            ("Regina González", "rgonzalez@vwys.com.mx"),
            ("Arturo Hernández", "ahernandez@vwys.com.mx"),
        ),
        tags=("energy-natural-resources", "environmental", "fracking"),
    ),
)


def clean(value: str) -> str:
    return SPACE_RE.sub(" ", value or "").strip()


def normalized(value: str) -> str:
    return clean(value).casefold().strip(" .:")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_pdf(path: Path) -> None:
    data = path.read_bytes()
    if len(data) < 1_000 or not data.startswith(b"%PDF-") or b"%%EOF" not in data[-2_048:]:
        raise ValueError(f"Invalid PDF source: {path}")


def linkify(value: str) -> str:
    output: list[str] = []
    cursor = 0
    for match in URL_RE.finditer(value):
        raw = match.group(0)
        url = raw.rstrip(".,);]")
        trailing = raw[len(url) :]
        output.append(html.escape(value[cursor : match.start()]))
        output.append(
            f'<a href="{html.escape(url, quote=True)}" target="_blank" rel="noopener noreferrer">{html.escape(url)}</a>'
        )
        output.append(html.escape(trailing))
        cursor = match.end()
    output.append(html.escape(value[cursor:]))
    return "".join(output).replace("\n", "<br>")


def render_run(run: Run) -> str:
    value = linkify(run.text)
    if not value:
        return ""
    if run.bold:
        value = f"<strong>{value}</strong>"
    if run.italic:
        value = f"<em>{value}</em>"
    if run.underline:
        value = f"<u>{value}</u>"
    return value


def render_paragraph_inline(paragraph: Paragraph) -> str:
    pieces: list[str] = []
    for item in paragraph.iter_inner_content():
        if isinstance(item, Run):
            pieces.append(render_run(item))
        elif isinstance(item, Hyperlink):
            label = "".join(render_run(run) for run in item.runs)
            address = item.address or ""
            if address.startswith(("https://", "http://", "mailto:")):
                pieces.append(
                    f'<a href="{html.escape(address, quote=True)}" target="_blank" rel="noopener noreferrer">{label}</a>'
                )
            else:
                pieces.append(label)
    return "".join(pieces) or linkify(paragraph.text)


def looks_heading(paragraph: Paragraph) -> bool:
    value = clean(paragraph.text)
    style = (paragraph.style.name or "").casefold()
    if style.startswith("heading"):
        return True
    if len(value) > 125 or value.endswith((".", ";", ",", "?", "!")):
        return False
    if re.match(r"^(?:[A-ZÁÉÍÓÚÑ]\.|\d+[.)]|[IVXLCDM]+\.)\s+", value):
        return True
    words = value.split()
    return 1 <= len(words) <= 11 and "@" not in value and not re.search(r"\+?\d{2,}", value)


def render_paragraph(paragraph: Paragraph) -> str:
    value = clean(paragraph.text)
    if not value:
        return ""
    tag = "h2" if looks_heading(paragraph) else "p"
    return f"<{tag}>{render_paragraph_inline(paragraph)}</{tag}>"


def render_table(table: Table) -> str:
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [clean(cell.text) for cell in row.cells]
        if any(cells):
            rows.append(cells)
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    head = rows[0] + [""] * (width - len(rows[0]))
    body = [row + [""] * (width - len(row)) for row in rows[1:]]
    head_html = "".join(f"<th>{linkify(cell)}</th>" for cell in head)
    body_html = "".join(
        "<tr>" + "".join(f"<td>{linkify(cell)}</td>" for cell in row) + "</tr>" for row in body
    )
    return f"<table><thead><tr>{head_html}</tr></thead><tbody>{body_html}</tbody></table>"


def iter_blocks(document: Document) -> Iterable[Paragraph | Table]:
    paragraph_by_element = {paragraph._p: paragraph for paragraph in document.paragraphs}
    table_by_element = {table._tbl: table for table in document.tables}
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield paragraph_by_element[child]
        elif isinstance(child, CT_Tbl):
            yield table_by_element[child]


def should_stop(value: str, authors: tuple[tuple[str, str], ...]) -> bool:
    current = normalized(value)
    if not current:
        return False
    if current == "***" or any(current.startswith(prefix) for prefix in STOP_PREFIXES):
        return True
    author_names = {normalized(name) for name, _ in authors}
    author_emails = {email.casefold() for _, email in authors}
    return (
        any(current == name or current.startswith(f"{name},") for name in author_names)
        or any(email and email in current for email in author_emails)
    )


def docx_content(
    path: Path,
    excerpt_ordinal: int,
    authors: tuple[tuple[str, str], ...],
    excerpt_override: str | None,
) -> tuple[str, str]:
    document = Document(path)
    meaningful = [paragraph for paragraph in document.paragraphs if clean(paragraph.text)]
    if excerpt_ordinal >= len(meaningful):
        raise ValueError(f"Excerpt paragraph not found in {path}")
    excerpt_paragraph = meaningful[excerpt_ordinal]
    excerpt_text = excerpt_override or clean(excerpt_paragraph.text)
    if len(excerpt_text) > 470:
        raise ValueError(f"Excerpt exceeds CMS limit in {path}: {len(excerpt_text)}")

    include_excerpt_in_body = bool(excerpt_override)
    started = False
    pieces: list[str] = []
    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            if block._p is excerpt_paragraph._p:
                started = True
                if include_excerpt_in_body:
                    pieces.append(render_paragraph(block))
                continue
            if not started:
                continue
            plain = clean(block.text)
            if should_stop(plain, authors):
                break
            rendered = render_paragraph(block)
        else:
            if not started:
                continue
            rendered = render_table(block)
        if rendered:
            pieces.append(rendered)

    content = "".join(pieces)
    if not content:
        raise ValueError(f"No body content extracted from {path}")
    return f"<p>{html.escape(excerpt_text)}</p>", content


def manual_content(paragraphs: tuple[str, ...]) -> tuple[str, str]:
    if len(paragraphs) < 2:
        raise ValueError("Manual source must contain an excerpt and body")
    excerpt = paragraphs[0]
    if len(excerpt) > 470:
        raise ValueError("Manual excerpt exceeds CMS limit")
    return f"<p>{html.escape(excerpt)}</p>", "".join(f"<p>{html.escape(value)}</p>" for value in paragraphs[1:])


def append_pdf(content: str, pdf_path: str, language: str) -> str:
    label = "Download the source note in PDF" if language == "en" else "Descargar la nota fuente en PDF"
    return content + f'<p><a href="{html.escape(pdf_path, quote=True)}" target="_blank" rel="noopener noreferrer">{label}</a></p>'


def build_note(config: NoteConfig, source_root: Path, asset_root: Path | None) -> dict[str, Any]:
    folder = source_root / config.folder
    pdf = folder / config.pdf
    pdf_es = folder / config.pdf_es
    for source_pdf in (pdf, pdf_es):
        validate_pdf(source_pdf)

    source_files: list[dict[str, Any]] = []
    if config.en and config.es:
        en_path = folder / config.en.path
        es_path = folder / config.es.path
        excerpt, content = docx_content(en_path, config.en.excerpt_ordinal, config.authors, config.excerpt)
        excerpt_es, content_es = docx_content(es_path, config.es.excerpt_ordinal, config.authors, config.excerpt_es)
        source_files.extend(
            [
                {"language": "en", "name": config.en.path, "sha256": sha256(en_path)},
                {"language": "es", "name": config.es.path, "sha256": sha256(es_path)},
            ]
        )
    else:
        excerpt, content = manual_content(config.manual_en)
        excerpt_es, content_es = manual_content(config.manual_es)
        docx_sources = sorted(folder.glob("*.docx"))
        source_files.extend(
            {"language": "bilingual-layout-source", "name": source.name, "sha256": sha256(source)}
            for source in docx_sources
        )

    public_pdf = f"{PUBLIC_PDF_PREFIX}/{config.pdf}"
    public_pdf_es = f"{PUBLIC_PDF_PREFIX}/{config.pdf_es}"
    content = append_pdf(content, public_pdf, "en")
    content_es = append_pdf(content_es, public_pdf_es, "es")
    for value in (excerpt, excerpt_es, content, content_es):
        if UNSAFE_HTML_RE.search(value):
            raise ValueError(f"Unsafe generated HTML in {config.key}")

    if asset_root:
        asset_root.mkdir(parents=True, exist_ok=True)
        shutil.copy2(pdf, asset_root / config.pdf)
        shutil.copy2(pdf_es, asset_root / config.pdf_es)

    record: dict[str, Any] = {
        "key": config.key,
        "sourceFolder": config.folder,
        "date": config.date,
        "slug": config.slug,
        "title": config.title,
        "titleEs": config.title_es,
        "excerpt": excerpt,
        "excerptEs": excerpt_es,
        "content": content,
        "contentEs": content_es,
        "category": "news",
        "categoryEs": "Noticias",
        "tags": list(config.tags),
        "published": True,
        "featuredHome": False,
        "pdf": {"path": public_pdf, "sha256": sha256(pdf), "size": pdf.stat().st_size},
        "pdfEs": {"path": public_pdf_es, "sha256": sha256(pdf_es), "size": pdf_es.stat().st_size},
        "authors": [{"name": name, "email": email} for name, email in config.authors],
        "sourceFiles": source_files,
    }
    digest_source = json.dumps(record, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    record["contentSha256"] = hashlib.sha256(digest_source).hexdigest()
    return record


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--asset-root", type=Path)
    args = parser.parse_args()

    items = [build_note(note, args.source_root, args.asset_root) for note in NOTES]
    if len(items) != 11 or len({item["slug"] for item in items}) != 11:
        raise ValueError("Expected eleven unique canonical Dropbox notes")
    snapshot: dict[str, Any] = {
        "snapshotDate": "2026-08-14",
        "source": "Dropbox — Notas informativas 2026 supplied by Von Wobeser y Sierra",
        "items": items,
    }
    snapshot_digest = json.dumps(snapshot, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    snapshot["snapshotSha256"] = hashlib.sha256(snapshot_digest).hexdigest()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Generated {len(items)} bilingual notes: {args.output}")


if __name__ == "__main__":
    main()
